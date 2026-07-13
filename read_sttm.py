import pathlib
from openpyxl import load_workbook

path = pathlib.Path(r'C:/Users/p.samuel/Documents/VS_workspace/STTM_B5_PROJECT.xlsx')
print('PATH:', path)
print('EXISTS:', path.exists())
wb = load_workbook(filename=path, read_only=True)
print('SHEETS:' + '|'.join(wb.sheetnames))

name = 'DATE_DIM'
if name not in wb.sheetnames:
    print('MISSING SHEET:', name)
else:
    ws = wb[name]
    headers = [str(cell).strip() if cell is not None else '' for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    print('SHEET:' + name + '|HEADER:' + ','.join(headers))
    for i, row in enumerate(ws.iter_rows(min_row=2, max_row=20, values_only=True), start=2):
        print('ROW', i, '|', [str(cell).strip() if cell is not None else '' for cell in row])
